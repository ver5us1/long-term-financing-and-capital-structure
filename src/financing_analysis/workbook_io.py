"""Read the source workbook or its portable, metadata-free JSON snapshot.

The committed snapshot keeps inputs, formulas, and Excel's saved results. It
allows the project to run without the original thesis files or Microsoft Excel.
The importer reads the workbook twice and never writes to the source file.
"""

from __future__ import annotations

import hashlib
import json
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

DEFAULT_SNAPSHOT = Path(__file__).resolve().parents[2] / "data/source/workbook_snapshot.json"


def _json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    raise TypeError(f"Unsupported workbook value: {type(value).__name__}")


def snapshot_from_xlsx(path: str | Path) -> dict[str, Any]:
    """Extract occupied cells with formula text and cached Excel values.

    Workbook author names, comments, and other document properties are excluded.
    ISO date/time strings preserve the original observation dates in JSON.
    """
    from openpyxl import load_workbook

    source = Path(path)
    workbook = load_workbook(source, data_only=False, read_only=True, keep_links=False)
    cached = load_workbook(source, data_only=True, read_only=True, keep_links=False)
    sheets = []
    try:
        for worksheet in workbook.worksheets:
            cells = {}
            cached_rows = cached[worksheet.title].iter_rows()
            for row, cache_row in zip(worksheet.iter_rows(), cached_rows):
                for cell, cache_cell in zip(row, cache_row):
                    if cell.value is None:
                        continue
                    if cell.data_type == "f":
                        cells[cell.coordinate] = {
                            "value": None,
                            "formula": cell.value,
                            "cached": _json_value(cache_cell.value),
                        }
                    else:
                        cells[cell.coordinate] = {"value": _json_value(cell.value)}
            sheets.append({"name": worksheet.title, "cells": cells})
    finally:
        workbook.close()
        cached.close()
    return {
        "schema_version": 1,
        "source": {
            "filename": source.name,
            "sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        },
        "sheets": sheets,
    }


def save_snapshot(snapshot: dict[str, Any], path: str | Path) -> Path:
    """Save the source evidence as UTF-8 JSON, preserving numerical precision."""
    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(
        json.dumps(snapshot, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    return destination


def load_snapshot(path: str | Path | None = None) -> dict[str, Any]:
    """Load schema version 1; use the repository's committed source by default."""
    source = Path(path) if path is not None else DEFAULT_SNAPSHOT
    snapshot = json.loads(source.read_text(encoding="utf-8"))
    if snapshot.get("schema_version") != 1 or not isinstance(snapshot.get("sheets"), list):
        raise ValueError(f"Unsupported workbook snapshot schema: {source}")
    return snapshot
